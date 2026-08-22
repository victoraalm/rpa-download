"""Testes da gravação em planilha, incluindo ida e volta pelo Excel real.

Uso:  python tests/spreadsheet_test.py

O teste que usa o Excel é pulado (não falha) numa máquina sem Excel instalado.
"""

from __future__ import annotations

import datetime as dt
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

from app import spreadsheet as sp  # noqa: E402

_falhas = []


def check(label, cond):
    print(f"  [{'OK ' if cond else 'FALHOU'}] {label}")
    if not cond:
        _falhas.append(label)


# ------------------------------------------------------------------ conversão
def test_numero_br():
    print("== Número pt-BR vindo como texto ==")
    check("-1.506,61 vira número", sp.numero_br("-1.506,61") == -1506.61)
    check("559,33 vira número", sp.numero_br("559,33") == 559.33)
    check("1.234.567,89 vira número", sp.numero_br("1.234.567,89") == 1234567.89)
    check("-0,01 vira número", sp.numero_br("-0,01") == -0.01)
    # Sem vírgula decimal continua texto: é código, não valor.
    check("'5.100' continua texto", sp.numero_br("5.100") == "5.100")
    check("zeros à esquerda preservados", sp.numero_br("000638") == "000638")
    check("'1400000094' continua texto", sp.numero_br("1400000094") == "1400000094")
    check("'BRL' intacto", sp.numero_br("BRL") == "BRL")
    check("None intacto", sp.numero_br(None) is None)
    check("número já numérico intacto", sp.numero_br(12.5) == 12.5)


def test_valor_para_excel():
    print("== Coerção de tipos para a ponte COM ==")
    check("time vira texto", sp.valor_para_excel(dt.time(14, 36, 25)) == "14:36:25")
    check("date vira datetime",
          sp.valor_para_excel(dt.date(2026, 8, 21)) == dt.datetime(2026, 8, 21))
    check("str intacto", sp.valor_para_excel("abc") == "abc")
    check("None intacto", sp.valor_para_excel(None) is None)


def test_cabecalho():
    print("== Conferência de cabeçalho ==")
    i, t, d = sp.comparar_cabecalho(["Conta", "Nome fornecedor"], ["Conta", "Nome do fornecedor"])
    check("rótulo escrito diferente ainda casa", i == 2 and not d)
    i, t, d = sp.comparar_cabecalho(["Conta", "Valor"], ["Valor", "Conta"])
    check("ordem trocada é divergência", i == 0 and len(d) == 2)
    i, t, d = sp.comparar_cabecalho(["A", "B"], ["A", "B", "C"])
    check("coluna a mais conta no total", t == 3 and i == 2)


def test_largura():
    print("== Ajuste de largura ==")
    check("recorta", sp.ajustar_largura([(1, 2, 3)], 2) == [(1, 2)])
    check("completa com None", sp.ajustar_largura([(1,)], 3) == [(1, None, None)])


def test_converter_colunas():
    print("== Conversão por coluna ==")
    linhas = [("x", "-1.506,61", "000638")]
    out = sp.converter_numeros_br(linhas, [2])
    check("só a coluna pedida é convertida",
          out == [("x", -1506.61, "000638")])
    check("sem colunas, nada muda", sp.converter_numeros_br(linhas, []) == linhas)


# -------------------------------------------------------------------- guardas
def test_planilha_aberta(d):
    print("== Guarda: planilha em uso ==")
    p = os.path.join(d, "livre.xlsx")
    openpyxl.Workbook().save(p)
    check("arquivo livre não acusa", sp.esta_aberta(p) is False)
    check("arquivo inexistente não acusa", sp.esta_aberta(os.path.join(d, "nao.xlsx")) is False)
    # Simula o arquivo de bloqueio que o Excel cria ao abrir.
    lock = os.path.join(d, "~$livre.xlsx")
    with open(lock, "w") as f:
        f.write("")
    check("arquivo de bloqueio do Excel acusa", sp.esta_aberta(p) is True)
    os.remove(lock)


def test_backup(d):
    print("== Backup com retenção ==")
    p = os.path.join(d, "alvo.xlsx")
    openpyxl.Workbook().save(p)
    pasta = os.path.join(d, "bk")
    copias = [sp.fazer_backup(p, pasta, manter=3) for _ in range(5)]
    restantes = os.listdir(pasta)
    check("backup criado", os.path.isfile(copias[-1]))
    check("retenção mantém só 3", len(restantes) == 3)


def test_localizar(d):
    print("== Localizar o arquivo baixado ==")
    pasta = os.path.join(d, "dl")
    os.makedirs(pasta, exist_ok=True)
    for nome in ["Base A.xlsx", "Base B.xlsx", "outro.csv"]:
        openpyxl.Workbook().save(os.path.join(pasta, nome)) if nome.endswith("xlsx") \
            else open(os.path.join(pasta, nome), "w").close()
    achado = sp.localizar_arquivo(pasta, "Base*.xlsx")
    check("acha pelo padrão", os.path.basename(achado).startswith("Base"))
    check("padrão vazio devolve vazio", sp.localizar_arquivo(pasta, "") == "")


# ----------------------------------------------------------- ida e volta real
def _tem_excel() -> bool:
    try:
        import win32com.client as w
        x = w.DispatchEx("Excel.Application")
        x.Quit()
        return True
    except Exception:  # noqa: BLE001
        return False


def _montar_destino(caminho):
    """Planilha de destino com cabeçalho na linha 1 e uma Tabela do Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    for c, nome in enumerate(["Conta", "Nome do fornecedor", "Montante"], start=1):
        ws.cell(row=1, column=c, value=nome)
    ws.cell(row=2, column=1, value="antigo")
    t = Table(displayName="TB", ref="A1:C2")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    wb.save(caminho)


def _montar_baixado(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, nome in enumerate(["Conta", "Nome fornecedor", "Montante"], start=1):
        ws.cell(row=1, column=c, value=nome)
    dados = [("271374", "ACME LTDA", "-1.506,61"),
             ("259181", "BETA SA", "90.129,82"),
             ("271374", "GAMA ME", "559,33")]
    for r, linha in enumerate(dados, start=2):
        for c, v in enumerate(linha, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(caminho)


def test_gravacao_real(d):
    print("== Gravação real via Excel ==")
    if not _tem_excel():
        print("  [PULADO] Excel não disponível nesta máquina")
        return
    destino = os.path.join(d, "indicador.xlsx")
    baixado = os.path.join(d, "Base Custos.xlsx")
    _montar_destino(destino)
    _montar_baixado(baixado)

    t = sp.Transferencia(
        arquivo="Base*.xlsx", planilha=destino, aba="Dados", tabela="TB",
        linha_inicial=2, colunas=3, linha_destino=2, coluna_destino=1,
        colunas_numero_br=[3], backups=5,
    )
    try:
        res = sp.gravar(t, baixado, os.path.join(d, "bk"))
    except sp.ErroPlanilha as e:
        if "SOMENTE LEITURA" in str(e):
            # Nesta maquina o Excel esta abrindo tudo em somente leitura (Office
            # em periodo de carencia). O que importa testar aqui e que a guarda
            # DISPARA em vez de reportar sucesso sem gravar nada.
            check("guarda de somente leitura dispara em vez de falhar calado",
                  "Nada foi gravado" in str(e))
            print("  [PULADO] round-trip: o Excel desta máquina abre em somente leitura")
            return
        raise
    check("gravou", res.ok and res.linhas == 3)
    check("backup registrado", os.path.isfile(res.backup))
    check("cabeçalho conferido", res.cabecalho_iguais == 3)

    wb = openpyxl.load_workbook(destino, data_only=True)
    ws = wb["Dados"]
    check("linha 1 chegou", ws.cell(row=2, column=1).value in ("271374", 271374))
    check("texto pt-BR virou NÚMERO",
          isinstance(ws.cell(row=2, column=3).value, (int, float))
          and abs(ws.cell(row=2, column=3).value + 1506.61) < 0.01)
    check("3ª linha chegou", ws.cell(row=4, column=2).value == "GAMA ME")
    check("tabela redimensionada", ws.tables["TB"].ref.endswith("4"))
    wb.close()


def test_recusa_layout_diferente(d):
    print("== Recusa layout embaralhado ==")
    if not _tem_excel():
        print("  [PULADO] Excel não disponível nesta máquina")
        return
    destino = os.path.join(d, "ind2.xlsx")
    baixado = os.path.join(d, "Trocado.xlsx")
    _montar_destino(destino)
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, nome in enumerate(["Montante", "Conta", "Nome"], start=1):  # ordem trocada
        ws.cell(row=1, column=c, value=nome)
    ws.cell(row=2, column=1, value="1,00")
    wb.save(baixado)

    t = sp.Transferencia(planilha=destino, aba="Dados", linha_inicial=2, colunas=3,
                         conferir_cabecalho=True)
    try:
        sp.gravar(t, baixado, os.path.join(d, "bk"))
        check("deveria ter recusado", False)
    except sp.ErroPlanilha as e:
        check("recusa layout embaralhado ou somente leitura",
              "não corresponde" in str(e) or "SOMENTE LEITURA" in str(e))


def test_recusa_planilha_aberta(d):
    print("== Recusa planilha aberta ==")
    destino = os.path.join(d, "ind3.xlsx")
    baixado = os.path.join(d, "Base Custos.xlsx")
    _montar_destino(destino)
    if not os.path.isfile(baixado):
        _montar_baixado(baixado)
    lock = os.path.join(d, "~$ind3.xlsx")
    with open(lock, "w") as f:
        f.write("")
    t = sp.Transferencia(planilha=destino, aba="Dados", colunas=3)
    try:
        sp.gravar(t, baixado)
        check("deveria ter recusado", False)
    except sp.ErroPlanilha as e:
        check("recusa com mensagem clara", "aberta no Excel" in str(e))
    finally:
        os.remove(lock)


def test_recusa_destino_inexistente(d):
    print("== Recusa destino inexistente ==")
    t = sp.Transferencia(planilha=os.path.join(d, "nao_existe.xlsx"), aba="Dados")
    try:
        sp.gravar(t, os.path.join(d, "Base Custos.xlsx"))
        check("deveria ter recusado", False)
    except sp.ErroPlanilha as e:
        check("explica por que não cria do zero", "não cria planilha" in str(e))


def main():
    with tempfile.TemporaryDirectory() as d:
        test_numero_br()
        test_valor_para_excel()
        test_cabecalho()
        test_largura()
        test_converter_colunas()
        test_planilha_aberta(d)
        test_backup(d)
        test_localizar(d)
        test_gravacao_real(d)
        test_recusa_layout_diferente(d)
        test_recusa_planilha_aberta(d)
        test_recusa_destino_inexistente(d)
    print()
    if _falhas:
        print(f"{len(_falhas)} verificação(ões) falharam:")
        for f in _falhas:
            print("  - " + f)
        return 1
    print("planilha: todas as verificações passaram")
    return 0


if __name__ == "__main__":
    sys.exit(main())
